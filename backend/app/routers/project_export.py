"""Полная выгрузка проекта одним архивом — для администратора.

Нужна, чтобы забрать всё из системы (например, перед закрытием хостинга):
файлы документов в структуре папок проекта и реестр с документами,
ревизиями, замечаниями и ответами.

Структура архива:
    IMP_выгрузка_2026-10-21/
      Реестр_IMP.xlsx          — Документы · Ревизии · Замечания и ответы · Файлы
      README.txt
      PD/
        IMP-...-AR33.2 — Название/
          Ревизия A (IFR)/
            IMP-...-AR33.2_A.pdf
            Редактируемые файлы/…
            Вложения к замечаниям/IMP-RMK-000012_файл.pdf
          Файлы документа/…        — вложения без привязки к ревизии
          Вложенные/
            IMP-...-AR33.2-01 — Название/…

Файлы кладутся без сжатия (ZIP_STORED): это PDF и сканы, которые почти не
жмутся, а архив на гигабайт собирается за секунды и укладывается в таймаут
прокси. Временный файл удаляется после отдачи.
"""

from __future__ import annotations

import os
import re
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from starlette.background import BackgroundTask

from app.database import get_db
from app.deps import get_current_user
from app.models import (
    Comment,
    CommentAttachment,
    Document,
    DocumentAttachment,
    MDRRecord,
    Project,
    Revision,
    User,
)
from app.routers.documents import (
    _EXPORT_STATUS_LABELS,
    _write_comments_sheet,
    _write_registry_sheets,
    list_documents_registry,
)

router = APIRouter()

_FORBIDDEN = re.compile(r'[<>:"/\\|?*\x00-\x1f]+')


def _safe(name: str, limit: int = 90) -> str:
    """Имя папки/файла, допустимое в Windows и macOS."""
    cleaned = _FORBIDDEN.sub("_", (name or "").strip())
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" .")
    if len(cleaned) > limit:
        cleaned = cleaned[:limit].rstrip(" .")
    return cleaned or "_"


def _safe_file(name: str, limit: int = 120) -> str:
    """Имя файла: режем основу, расширение сохраняем."""
    path = Path(_safe(name, limit=255))
    suffix = path.suffix[:12]
    stem = path.stem[: max(1, limit - len(suffix))]
    return f"{stem}{suffix}"


@router.get("/projects/{project_id}/export.zip")
def export_project_archive(
    project_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if current_user.role.value != "admin":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Полная выгрузка проекта доступна только администратору",
        )
    project = db.query(Project).filter(Project.id == project_id).first()
    if project is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Проект не найден")

    stamp = datetime.utcnow().strftime("%Y-%m-%d")
    root = _safe(f"{project.code}_выгрузка_{stamp}")

    mdrs = db.query(MDRRecord).filter(MDRRecord.project_code == project.code).all()
    mdr_by_id = {m.id: m for m in mdrs}
    documents = db.query(Document).filter(Document.mdr_id.in_(list(mdr_by_id) or [-1])).all()
    document_by_mdr = {d.mdr_id: d for d in documents}

    # --- папка каждого документа: категория / головной / Вложенные / вложенный
    def top_folder(mdr: MDRRecord) -> str:
        document = document_by_mdr.get(mdr.id)
        number = document.document_num if document else mdr.doc_number
        return _safe(f"{number} — {mdr.doc_name}")

    folder_by_mdr: dict[int, str] = {}
    for mdr in mdrs:
        category = _safe(mdr.category or "Без категории", limit=40)
        parent = mdr_by_id.get(mdr.parent_id) if mdr.parent_id else None
        if parent is not None:
            folder_by_mdr[mdr.id] = f"{root}/{_safe(parent.category or 'Без категории', limit=40)}/{top_folder(parent)}/Вложенные/{top_folder(mdr)}"
        else:
            folder_by_mdr[mdr.id] = f"{root}/{category}/{top_folder(mdr)}"
    folder_by_document = {d.id: folder_by_mdr[d.mdr_id] for d in documents if d.mdr_id in folder_by_mdr}

    revisions = (
        db.query(Revision)
        .filter(Revision.document_id.in_(list(folder_by_document) or [-1]))
        .order_by(Revision.id.asc())
        .all()
    )
    revision_folder: dict[int, str] = {}
    seen_revision_folders: dict[str, int] = {}
    for revision in revisions:
        purpose = f" ({revision.issue_purpose})" if revision.issue_purpose else ""
        folder = f"{folder_by_document[revision.document_id]}/{_safe(f'Ревизия {revision.revision_code}{purpose}', limit=40)}"
        # Тот же код ревизии бывает дважды (перевыпуск после RJ) — иначе
        # файлы двух выпусков смешались бы в одной папке.
        seen_revision_folders[folder] = seen_revision_folders.get(folder, 0) + 1
        if seen_revision_folders[folder] > 1:
            folder = f"{folder} — выпуск {seen_revision_folders[folder]}"
        revision_folder[revision.id] = folder

    # --- список файлов: (путь на диске, путь в архиве, описание для листа «Файлы»)
    planned: list[tuple[str | None, str, dict]] = []
    document_by_id = {d.id: d for d in documents}
    for revision in revisions:
        if not revision.file_path:
            continue
        document = document_by_id[revision.document_id]
        planned.append((
            revision.file_path,
            f"{revision_folder[revision.id]}/{_safe_file(f'{document.document_num}_{revision.revision_code}.pdf')}",
            {"kind": "PDF ревизии", "document": document.document_num, "revision": revision.revision_code},
        ))
    for item in db.query(DocumentAttachment).filter(DocumentAttachment.document_id.in_(list(document_by_id) or [-1])).all():
        document = document_by_id[item.document_id]
        if item.revision_id and item.revision_id in revision_folder:
            folder = f"{revision_folder[item.revision_id]}/Редактируемые файлы"
            revision_code = next((r.revision_code for r in revisions if r.id == item.revision_id), "")
        else:
            folder = f"{folder_by_document[item.document_id]}/Файлы документа"
            revision_code = ""
        planned.append((
            item.file_path,
            f"{folder}/{_safe_file(item.file_name)}",
            {"kind": "Файл документа", "document": document.document_num, "revision": revision_code},
        ))
    comment_rows = (
        db.query(CommentAttachment, Comment)
        .join(Comment, Comment.id == CommentAttachment.comment_id)
        .filter(Comment.revision_id.in_(list(revision_folder) or [-1]))
        .all()
    )
    revision_by_id = {r.id: r for r in revisions}
    for attachment, comment in comment_rows:
        revision = revision_by_id[comment.revision_id]
        label = comment.remark_number or f"comment-{comment.id}"
        planned.append((
            attachment.file_path,
            f"{revision_folder[revision.id]}/Вложения к замечаниям/{_safe_file(f'{label}_{attachment.file_name}')}",
            {
                "kind": "Вложение к замечанию",
                "document": document_by_id[revision.document_id].document_num,
                "revision": revision.revision_code,
                "remark": label,
            },
        ))

    # --- реестр
    from openpyxl import Workbook
    from openpyxl.styles import Font

    registry_rows = list_documents_registry(
        project_code=project.code,
        category=None,
        discipline_code=None,
        document_title=None,
        release_status=None,
        revision_status=None,
        comments_scope=None,
        overdue_only=False,
        for_reporting=True,
        db=db,
        current_user=current_user,
    )
    pdf_in_archive = {}
    for disk_path, arc_path, meta in planned:
        if meta["kind"] == "PDF ревизии":
            pdf_in_archive.setdefault((meta["document"], meta["revision"]), arc_path[len(root) + 1:])
    parent_number = {
        d.id: (
            document_by_mdr[mdr_by_id[d.mdr_id].parent_id].document_num
            if mdr_by_id[d.mdr_id].parent_id and mdr_by_id[d.mdr_id].parent_id in document_by_mdr
            else ""
        )
        for d in documents
    }
    wb = Workbook()
    _write_registry_sheets(
        wb,
        registry_rows,
        extra_doc_columns=["Головной документ", "Папка в архиве"],
        extra_doc_values={
            row.document_id: [parent_number.get(row.document_id, ""), folder_by_document.get(row.document_id, "")[len(root) + 1:]]
            for row in registry_rows
        },
        extra_rev_columns=["PDF в архиве"],
        extra_rev_values={
            revision.id: [pdf_in_archive.get((row.document_num, revision.revision_code), "нет PDF")]
            for row in registry_rows
            for revision in row.revisions
        },
    )
    remarks_ws = wb.create_sheet("Замечания и ответы")
    records = (
        db.query(Comment, Revision, Document, MDRRecord, User)
        .join(Revision, Revision.id == Comment.revision_id)
        .join(Document, Document.id == Revision.document_id)
        .join(MDRRecord, MDRRecord.id == Document.mdr_id)
        .outerjoin(User, User.id == Comment.author_id)
        .filter(MDRRecord.project_code == project.code)
        .order_by(Document.document_num, Revision.id, Comment.id)
        .all()
    )
    _write_comments_sheet(db, remarks_ws, records)

    files_ws = wb.create_sheet("Файлы")
    files_ws.append(["Тип", "Документ", "Ревизия", "Замечание", "Путь в архиве", "Состояние"])
    for cell in files_ws[1]:
        cell.font = Font(bold=True)

    # --- архив
    temp = tempfile.NamedTemporaryFile(prefix=f"project_{project.code}_", suffix=".zip", delete=False)
    temp_path = temp.name
    temp.close()
    used: set[str] = set()
    packed = missing = 0

    def unique(arc_path: str) -> str:
        candidate, index = arc_path, 1
        while candidate in used:
            stem, ext = os.path.splitext(arc_path)
            candidate = f"{stem}_{index}{ext}"
            index += 1
        used.add(candidate)
        return candidate

    with zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_STORED) as archive:
        # Папка каждого документа и каждой ревизии — даже без файлов: архив
        # повторяет всю структуру проекта, а не только то, что загружено.
        for folder in sorted(set(folder_by_document.values()) | set(revision_folder.values())):
            archive.writestr(f"{folder}/", "")
        for disk_path, arc_path, meta in planned:
            arc_name = unique(arc_path)
            if disk_path and Path(disk_path).is_file():
                archive.write(disk_path, arcname=arc_name)
                packed += 1
                state = "в архиве"
            else:
                missing += 1
                state = "файл не найден на сервере"
            files_ws.append([
                meta["kind"], meta.get("document", ""), meta.get("revision", ""), meta.get("remark", ""),
                arc_name[len(root) + 1:], state,
            ])
        for index, width in enumerate([22, 32, 10, 18, 90, 26], start=1):
            files_ws.column_dimensions[files_ws.cell(row=1, column=index).column_letter].width = width

        registry_tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        registry_tmp.close()
        try:
            wb.save(registry_tmp.name)
            archive.write(registry_tmp.name, arcname=f"{root}/Реестр_{_safe(project.code, 20)}.xlsx")
        finally:
            os.unlink(registry_tmp.name)

        readme = "\n".join([
            f"Выгрузка проекта {project.code} — {project.name}",
            f"Сформирована: {datetime.utcnow().strftime('%d.%m.%Y %H:%M')} UTC, пользователь {current_user.email}",
            "",
            f"Документов: {len(documents)}, ревизий: {len(revisions)}, замечаний и ответов: {len(records)}",
            f"Файлов в архиве: {packed}" + (f", не найдено на сервере: {missing} (см. лист «Файлы» в реестре)" if missing else ""),
            "",
            "Структура: категория / документ / ревизия. Вложенные документы — в папке «Вложенные»",
            "внутри головного. Редактируемые файлы ревизии — «Редактируемые файлы», вложения",
            "к замечаниям — «Вложения к замечаниям» (имя начинается с номера замечания).",
            "",
            "Реестр_*.xlsx: листы «Документы», «Ревизии» (с путём к PDF в архиве),",
            "«Замечания и ответы» и «Файлы» (полный список с путями).",
            "",
            "Этапы рассмотрения: " + "; ".join(f"{v}" for v in _EXPORT_STATUS_LABELS.values()),
        ])
        archive.writestr(f"{root}/README.txt", readme)

    return FileResponse(
        path=temp_path,
        media_type="application/zip",
        filename=f"{root}.zip",
        background=BackgroundTask(os.unlink, temp_path),
    )
